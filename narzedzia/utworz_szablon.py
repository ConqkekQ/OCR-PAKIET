"""Tworzy pakiety.xlsx z przykładowymi pakietami (do zastąpienia prawdziwymi)."""
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

out = sys.argv[1] if len(sys.argv) > 1 else "pakiety.xlsx"
wb = Workbook()
ws = wb.active
ws.title = "Pakiety"
rows = [
    ("Pakiet", "Kod towaru", "Nazwa towaru", "Uwagi"),
    ("PRZYKŁAD Eficur 250 ml + 100 ml", "L008877", "Eficur 50mg/ml, 250 ml", "Przykład – usuń i wpisz prawdziwe pakiety"),
    ("PRZYKŁAD Eficur 250 ml + 100 ml", "LHPEFI100", "Eficur 50mg/ml, 100 ml", ""),
    ("PRZYKŁAD Pyralgivet duet", "LVAPYRG250", "Pyralgivet inj. 250 ml", "Przykład"),
    ("PRZYKŁAD Pyralgivet duet", "L002591", "Pyralgivet inj. 100 ml", ""),
    ("PRZYKŁAD Rękawice S/M/L", "MREKLATS100", "Rękawice lateksowe S 100 szt.", "Przykład pakietu z brakującym towarem"),
    ("PRZYKŁAD Rękawice S/M/L", "MREKLATM100", "Rękawice lateksowe M 100 szt.", ""),
    ("PRZYKŁAD Rękawice S/M/L", "MREKLATL100", "Rękawice lateksowe L 100 szt.", ""),
]
for r in rows:
    ws.append(r)
for c, w in zip("ABCD", (38, 16, 42, 50)):
    ws.column_dimensions[c].width = w
tab = Table(displayName="Pakiety", ref=f"A1:D{len(rows)}")
tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws.add_table(tab)
ws.freeze_panes = "A2"

wt = wb.create_sheet("Towary")
wt.append(("Kod towaru", "Nazwa towaru"))
wt.column_dimensions["A"].width = 18
wt.column_dimensions["B"].width = 60
for cell in wt[1]:
    cell.font = Font(bold=True)
    cell.fill = PatternFill("solid", fgColor="DDEBF7")
wt.freeze_panes = "A2"

wi = wb.create_sheet("Instrukcja")
lines = [
    "Jak uzupełnić plik",
    "",
    "Arkusz Pakiety: jeden wiersz = jeden towar w pakiecie.",
    "• Pakiet – nazwa pakietu; wszystkie towary tego samego pakietu mają identyczną nazwę.",
    "• Kod towaru – kod z kolumny Towar w XL (np. L008877). Wielkość liter bez znaczenia.",
    "• Nazwa towaru – opcjonalnie, wyświetla się w podpowiedzi.",
    "• Uwagi – opcjonalnie, np. warunek pakietu lub rabat; wystarczy w jednym wierszu pakietu.",
    "Ten sam towar może należeć do wielu pakietów.",
    "",
    "Arkusz Towary (opcjonalny, ale zalecany): pełna lista kodów towarów z XL.",
    "Dzięki niej nakładka lepiej odróżnia błędy odczytu od towarów spoza pakietów.",
    "Wystarczy wkleić kolumnę kodów z eksportu kartoteki towarów.",
    "",
    "Po zapisaniu pliku nakładka sama wczyta zmiany (lub kliknij ↻ Pakiety).",
]
for l in lines:
    wi.append((l,))
wi["A1"].font = Font(bold=True, size=13)
wi.column_dimensions["A"].width = 100
for row in wi.iter_rows():
    for c in row:
        c.alignment = Alignment(wrap_text=True)
wb.save(out)
print("Zapisano", out)
