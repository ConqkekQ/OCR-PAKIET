"""Definicje pakietów (Excel) i dopasowanie odczytanych kodów do pakietów."""
from __future__ import annotations

import os
from dataclasses import dataclass, field

from openpyxl import load_workbook


@dataclass
class Item:
    code: str
    name: str = ""


@dataclass
class Package:
    name: str
    items: list[Item] = field(default_factory=list)
    note: str = ""

    @property
    def codes(self) -> list[str]:
        return [i.code for i in self.items]


@dataclass
class Suggestion:
    package: Package
    present: list[Item]
    missing: list[Item]

    @property
    def complete(self) -> bool:
        return not self.missing


def _clean_code(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip().upper()


def _find_cols(header: list, wanted: dict[str, tuple]) -> dict[str, int]:
    norm = [str(h or "").strip().lower() for h in header]
    out = {}
    for key, names in wanted.items():
        for i, h in enumerate(norm):
            if any(h.startswith(n) for n in names):
                out[key] = i
                break
    return out


class PackageBook:
    """Pakiety z arkusza "Pakiety" + opcjonalna lista wszystkich towarów z arkusza "Towary"."""

    def __init__(self, path: str):
        self.path = path
        self.mtime = 0.0
        self.packages: list[Package] = []
        self.by_code: dict[str, list[Package]] = {}
        self.names: dict[str, str] = {}
        self.universe: set[str] = set()
        self.matcher = CodeMatcher(set(), set())
        self.load()

    def changed_on_disk(self) -> bool:
        try:
            return os.path.getmtime(self.path) != self.mtime
        except OSError:
            return False

    def load(self) -> None:
        if not os.path.exists(self.path):
            raise FileNotFoundError(f"Brak pliku z pakietami: {self.path}")
        wb = load_workbook(self.path, read_only=True, data_only=True)
        ws = wb["Pakiety"] if "Pakiety" in wb.sheetnames else wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            raise ValueError("Arkusz z pakietami jest pusty")
        cols = _find_cols(list(rows[0]), {
            "pakiet": ("pakiet",), "kod": ("kod",), "nazwa": ("nazwa",), "uwagi": ("uwag", "opis", "warun"),
        })
        if "pakiet" not in cols or "kod" not in cols:
            raise ValueError('W pierwszym wierszu arkusza muszą być kolumny "Pakiet" i "Kod towaru"')

        pk: dict[str, Package] = {}
        for r in rows[1:]:
            def cell(key):
                i = cols.get(key)
                return r[i] if i is not None and i < len(r) else None

            pname = str(cell("pakiet") or "").strip()
            code = _clean_code(cell("kod"))
            if not pname or not code:
                continue
            p = pk.setdefault(pname, Package(pname))
            name = str(cell("nazwa") or "").strip()
            if code not in p.codes:
                p.items.append(Item(code, name))
            if name:
                self.names.setdefault(code, name)
            note = str(cell("uwagi") or "").strip()
            if note and not p.note:
                p.note = note

        universe = set()
        if "Towary" in wb.sheetnames:
            trows = list(wb["Towary"].iter_rows(values_only=True))
            if trows:
                tcols = _find_cols(list(trows[0]), {"kod": ("kod",), "nazwa": ("nazwa",)})
                ki, ni = tcols.get("kod", 0), tcols.get("nazwa")
                for r in trows[1:]:
                    code = _clean_code(r[ki] if ki < len(r) else None)
                    if code:
                        universe.add(code)
                        if ni is not None and ni < len(r) and r[ni]:
                            self.names.setdefault(code, str(r[ni]).strip())
        wb.close()

        self.packages = list(pk.values())
        self.by_code = {}
        for p in self.packages:
            for c in p.codes:
                self.by_code.setdefault(c, []).append(p)
        self.universe = universe
        self.matcher = CodeMatcher(set(self.by_code), universe)
        self.mtime = os.path.getmtime(self.path)

    def suggest(self, order_codes: set[str]) -> list[Suggestion]:
        out = []
        for p in self.packages:
            present = [i for i in p.items if i.code in order_codes]
            if not present:
                continue
            missing = [i for i in p.items if i.code not in order_codes]
            out.append(Suggestion(p, present, missing))
        # najpierw pakiety niekompletne (jest co dopowiedzieć), te z największą liczbą trafień wyżej
        out.sort(key=lambda s: (s.complete, -len(s.present), len(s.missing), s.package.name))
        return out


# ---------------- dopasowanie kodów odczytanych przez OCR ----------------

_CANON = str.maketrans({"O": "0", "Q": "0", "I": "1", "|": "1", "S": "5", "Z": "2", "B": "8", "G": "6"})


def canon(code: str) -> str:
    return code.upper().translate(_CANON)


def levenshtein(a: str, b: str, limit: int = 3) -> int:
    if abs(len(a) - len(b)) > limit:
        return limit + 1
    prev = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        cur = [i]
        best = i
        for j, cb in enumerate(b, 1):
            v = min(prev[j] + 1, cur[j - 1] + 1, prev[j - 1] + (ca != cb))
            cur.append(v)
            best = min(best, v)
        if best > limit:
            return limit + 1
        prev = cur
    return prev[-1]


class CodeMatcher:
    """Mapuje kod odczytany przez OCR na prawdziwy kod towaru.

    1. dokładne trafienie,
    2. trafienie po ujednoliceniu mylonych znaków (O/0, I/1, S/5, Z/2, B/8, G/6),
    3. jeden błąd OCR (litera za dużo/za mało/inna) – tylko gdy w bazie nie ma
       innego podobnego kodu, który mógłby być prawdziwy.
    Znane kody = kody z pakietów + opcjonalny arkusz "Towary" (pełna lista towarów
    zwiększa pewność, bo pozwala odróżnić towary spoza pakietów).
    """

    def __init__(self, package_codes: set[str], universe: set[str]):
        self.known = set(package_codes) | set(universe)
        self.by_canon: dict[str, list[str]] = {}
        for c in self.known:
            self.by_canon.setdefault(canon(c), []).append(c)
        self._cache: dict[str, str | None] = {}

    def match(self, ocr_code: str) -> str | None:
        if ocr_code in self._cache:
            return self._cache[ocr_code]
        res = self._match(ocr_code)
        self._cache[ocr_code] = res
        return res

    def _match(self, code: str) -> str | None:
        code = code.upper()
        if code in self.known:
            return code
        cc = canon(code)
        hits = self.by_canon.get(cc, [])
        if len(hits) == 1:
            return hits[0]
        if len(cc) < 6:
            return None
        scored = []
        for kc, originals in self.by_canon.items():
            d = levenshtein(cc, kc, 2)
            if d <= 2:
                scored.append((d, originals))
        best = [s for s in scored if s[0] == 1]
        if len(best) == 1 and len(best[0][1]) == 1 and len(scored) == 1:
            target = best[0][1][0]
            if _plausible_ocr_slip(cc, canon(target)):
                return target
        return None


def _plausible_ocr_slip(a: str, b: str) -> bool:
    """Czy różnica jednego znaku wygląda na błąd OCR, a nie na inny towar?
    Cyfra zamieniona na inną cyfrę (L002594 vs L002591) to najczęściej inny towar,
    więc takiej różnicy nie poprawiamy. Dopuszczamy: dodatkową/brakującą kreskę
    (I/l/1, kropka, myślnik) oraz pomyloną literę."""
    if len(a) == len(b):
        diff = [(x, y) for x, y in zip(a, b) if x != y]
        if len(diff) != 1:
            return False
        x, y = diff[0]
        return not (x.isdigit() and y.isdigit())
    longer, shorter = (a, b) if len(a) > len(b) else (b, a)
    for i in range(len(longer)):
        if longer[:i] + longer[i + 1:] == shorter:
            return longer[i] in "1.-_/" or longer[i].isalpha()
    return False
